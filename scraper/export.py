"""Write the Excel workbook: Contacts | Needs Review | Run Log."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .validate import COLUMNS

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
MAX_WIDTH = 60


def _write_sheet(ws, columns, rows):
    ws.append(columns)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append([row.get(c) for c in columns])
    for idx, col in enumerate(columns, start=1):
        longest = max([len(str(col))] + [len(str(r.get(col) or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, MAX_WIDTH)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def export_workbook(path, contacts, review, runs):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    _write_sheet(ws, COLUMNS, contacts)
    _write_sheet(wb.create_sheet("Needs Review"), COLUMNS, review)
    run_cols = ["id", "started_at", "finished_at", "status", "requests", "listing_pages", "profiles", "note"]
    _write_sheet(wb.create_sheet("Run Log"), run_cols, runs)
    wb.save(path)
    return path
