"""Excel writer: Records | Needs Review | Run Log. Never overwrites an existing file."""

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from core.clean import build_rows, columns_for
from core.recipe import Recipe

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
MAX_WIDTH = 60
RUN_COLUMNS = ["id", "started_at", "finished_at", "status", "requests", "pages", "records", "note"]
DATA_NOTE = ("Personal data. Scrub phone numbers against the DND registry before calls/SMS, include an "
             "opt-out in every message, and delete a person's details on request.")


def _write_sheet(ws, columns, rows):
    ws.append(columns)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append([row.get(c) for c in columns])
    for idx, col in enumerate(columns, start=1):
        longest = max([len(str(col))] + [len(str(r.get(col) or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, MAX_WIDTH)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def _free_path(folder: Path, stem: str) -> Path:
    path = folder / f"{stem}.xlsx"
    n = 2
    while path.exists():
        path = folder / f"{stem}_{n}.xlsx"
        n += 1
    return path


def safe_stem(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_")[:40] or "recipe"


def export_workbook(folder, recipe: Recipe, records: list[dict], runs: list[dict]):
    """Write a new workbook; returns (path, ready_count, review_count)."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    ready, review = build_rows(recipe, records)
    columns = columns_for(recipe)

    wb = Workbook()
    ws = wb.active
    ws.title = "Records"
    _write_sheet(ws, columns, ready)
    _write_sheet(wb.create_sheet("Needs Review"), columns, review)
    log_ws = wb.create_sheet("Run Log")
    _write_sheet(log_ws, RUN_COLUMNS, runs)
    log_ws.append([])
    log_ws.append([DATA_NOTE])

    path = _free_path(folder, f"{safe_stem(recipe.name)}_{datetime.now():%Y%m%d_%H%M}")
    wb.save(path)
    return path, len(ready), len(review)
