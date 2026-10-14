"""Master list: several recipes' results in one workbook, duplicates removed across recipes.

Each recipe's rows are first cleaned and de-duplicated on their own (core.clean.build_rows). Then rows from
all recipes are compared by phone number (any phone-type field), or by email when there is no phone.
The first occurrence stays in "Master"; later ones go to "Needs Review" as "duplicate of <recipe> <id>".
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from core.clean import FIXED_AFTER, build_rows
from core.export import _free_path, _write_sheet
from core.recipe import Recipe

FIXED_BEFORE = ["Master ID", "Recipe", "Record ID"]
SOURCE_COLUMNS = ["Recipe", "Ready rows", "Need review", "Duplicates of other recipes", "Last run", "Last status"]


def _match_key(row, phone_fields, email_fields):
    for name in phone_fields:
        if row.get(name):
            return ("phone", row[name].split(",")[0].strip())
    for name in email_fields:
        if row.get(name):
            return ("email", row[name].split(",")[0].strip().lower())
    return None


def build_master(sources: list[tuple[Recipe, list[dict]]]):
    """sources: [(recipe, stored records)]. Returns (columns, master_rows, review_rows, per_recipe_counts)."""
    field_names, seen = [], {}
    master, review, counts = [], [], []
    for recipe, records in sources:
        for f in recipe.all_fields():
            if f.name not in field_names:
                field_names.append(f.name)
        ready, needs_review = build_rows(recipe, records)
        phones = [f.name for f in recipe.all_fields() if f.type == "phone"]
        emails = [f.name for f in recipe.all_fields() if f.type == "email"]
        duplicates = 0
        for row in ready:
            row["Recipe"] = recipe.name
            key = _match_key(row, phones, emails)
            if key and key in seen:
                first = seen[key]
                extra = f"duplicate of {first['Recipe']} {first['Record ID']}"
                row["Issues"] = f"{row['Issues']}; {extra}" if row.get("Issues") else extra
                review.append(row)
                duplicates += 1
            else:
                if key:
                    seen[key] = row
                master.append(row)
        for row in needs_review:
            row["Recipe"] = recipe.name
            review.append(row)
        counts.append({"Recipe": recipe.name, "Ready rows": len(ready) - duplicates,
                       "Need review": len(needs_review), "Duplicates of other recipes": duplicates})
    for n, row in enumerate(master, 1):
        row["Master ID"] = f"M-{n:05d}"
    columns = FIXED_BEFORE + field_names + FIXED_AFTER
    return columns, master, review, counts


def export_master(folder, sources: list[tuple[Recipe, list[dict], dict | None]]):
    """sources: [(recipe, records, last_run or None)]. Writes a new workbook; returns (path, master, review)."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    columns, master, review, counts = build_master([(r, recs) for r, recs, _run in sources])
    for count, (_r, _recs, run) in zip(counts, sources):
        count["Last run"] = (run or {}).get("started_at", "")
        count["Last status"] = (run or {}).get("status", "never run")

    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    _write_sheet(ws, columns, master)
    _write_sheet(wb.create_sheet("Needs Review"), columns, review)
    _write_sheet(wb.create_sheet("Sources"), SOURCE_COLUMNS, counts)
    path = _free_path(folder, f"Master_list_{datetime.now():%Y%m%d_%H%M}")
    wb.save(path)
    return path, len(master), len(review)
