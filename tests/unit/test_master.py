import tempfile
import unittest

from openpyxl import load_workbook

from core.master import build_master, export_master
from core.recipe import from_dict


def recipe(name, fields):
    return from_dict({"name": name, "start_url": "https://example.com/", "list": {"item_selector": "li", "fields": fields}})


SHOPS = recipe("Shops Pune", [{"name": "Name", "selector": "b", "required": True},
                              {"name": "Phone", "selector": "i", "type": "phone"}])
CLINICS = recipe("Clinics Pune", [{"name": "Name", "selector": "b", "required": True},
                                  {"name": "Mobile", "selector": "i", "type": "phone"},
                                  {"name": "Email", "selector": "u", "type": "email"}])


def rec(key, **data):
    return {"item_key": key, "collected_at": "2026-10-14T10:00:00", "data": {"_start_page": "https://example.com/", **data}}


class MasterTests(unittest.TestCase):
    def sources(self):
        return [
            (SHOPS, [rec("a1", Name="Alpha Stores", Phone="98450 00001"),
                     rec("a2", Name="Beta Mart", Phone="98450 00002"),
                     rec("a3", Name="No Phone Shop")]),
            (CLINICS, [rec("b1", Name="Alpha Clinic", Mobile="+91 98450-00001", Email="x@y.in"),
                       rec("b2", Name="Gamma Care", Email="CARE@gamma.in"),
                       rec("b3", Name=None, Mobile="98450 00009")]),
        ]

    def test_cross_recipe_duplicates_and_columns(self):
        columns, master, review, counts = build_master(self.sources())
        self.assertEqual(columns[:3], ["Master ID", "Recipe", "Record ID"])
        self.assertEqual(columns[3:7], ["Name", "Phone", "Mobile", "Email"])
        self.assertEqual(columns[-4:], ["Source URL", "Start Page", "Collected On", "Issues"])
        names = [r["Name"] for r in master]
        self.assertEqual(names, ["Alpha Stores", "Beta Mart", "No Phone Shop", "Gamma Care"])
        self.assertEqual([r["Master ID"] for r in master][-1], "M-00004")
        dup = next(r for r in review if r["Name"] == "Alpha Clinic")
        self.assertIn("duplicate of Shops Pune R-00001", dup["Issues"])
        missing = next(r for r in review if r["Recipe"] == "Clinics Pune" and r["Name"] is None)
        self.assertIn("Name: missing", missing["Issues"])
        self.assertEqual(counts[1], {"Recipe": "Clinics Pune", "Ready rows": 1, "Need review": 1,
                                     "Duplicates of other recipes": 1})

    def test_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            sources = [(r, recs, {"started_at": "2026-10-14T09:00:00", "status": "done"}) for r, recs in self.sources()]
            path, n_master, n_review = export_master(tmp, sources)
            self.assertEqual((n_master, n_review), (4, 2))
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Master", "Needs Review", "Sources"])
            self.assertEqual(wb["Sources"]["F2"].value, "done")
            self.assertTrue(path.name.startswith("Master_list_"))


if __name__ == "__main__":
    unittest.main()
