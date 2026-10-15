"""End-to-end runner tests against the local fixture site. Never touches real websites."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from core.export import export_workbook
from core.fetcher import Fetcher
from core.recipe import from_dict
from core.runner import RunControl, run_recipe
from core.store import Store
from tests.integration.fixture_site import FixtureSite


def make_recipe(base, path="/list", pagination=None):
    return from_dict({
        "name": f"fixture {path}",
        "start_url": base + path,
        "list": {"item_selector": "#results li.item", "fields": [
            {"name": "Name", "selector": "b.name", "required": True},
            {"name": "Profile", "selector": "a.more", "attr": "href", "type": "link"},
        ]},
        "pagination": pagination or {"type": "page_param", "param": "page", "start": 1},
        "detail": {"link_field": "Profile", "fields": [
            {"name": "Phone", "label": "Phone Number", "label_scope": "table.info tr", "type": "phone",
             "required": True},
            {"name": "Pincode", "label": "Address", "label_scope": "table.info tr", "type": "pincode"},
        ]},
        "dedup_keys": ["Name", "Phone"],
    })


class RunnerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.site = FixtureSite().__enter__()
        self.store = Store(Path(self.tmp.name) / "t.db", allow_local_urls=True)

    def tearDown(self):
        self.store.close()
        self.site.__exit__(None, None, None)
        self.tmp.cleanup()

    def fetcher(self):
        return Fetcher(self.store, min_delay=0, max_delay=0, retry_wait=0, max_retries=1, _min_delay_floor=0)

    def run_path(self, path="/list", **kwargs):
        pagination = kwargs.pop("pagination", None)
        rid = self.store.save_recipe(make_recipe(self.site.base, path, pagination))
        return rid, run_recipe(self.store, rid, self.fetcher(), **kwargs)

    def test_full_run_and_export(self):
        events = []
        rid, result = self.run_path(on_progress=events.append)
        self.assertEqual(result.status, "done")
        self.assertEqual(result.records, 6)
        self.assertEqual(result.pages, 3)  # page 3 is empty and ends the run
        self.assertTrue(any(e.event == "page" for e in events))

        path, ready, review = export_workbook(Path(self.tmp.name), self.store.get_recipe(rid),
                                              self.store.records(rid), self.store.runs(rid))
        self.assertEqual((ready, review), (4, 2))  # 5 has no phone, 6 duplicates 1
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Records", "Needs Review", "Run Log"])
        header = [c.value for c in wb["Records"][1]]
        self.assertEqual(header, ["Record ID", "Name", "Profile", "Phone", "Pincode", "Source URL",
                                  "Start Page", "Collected On", "Issues"])
        first = [c.value for c in wb["Records"][2]]
        self.assertEqual(first[1], "DR. ALPHA")
        self.assertEqual(first[6], self.site.base + "/list")
        self.assertEqual(first[3], "+919845000001")
        self.assertEqual(first[4], "560001")
        # never overwrite
        path2, _, _ = export_workbook(Path(self.tmp.name), self.store.get_recipe(rid),
                                      self.store.records(rid), self.store.runs(rid))
        self.assertNotEqual(path, path2)

    def test_deleted_person_is_not_collected_again(self):
        rid, result = self.run_path()
        self.assertEqual(result.records, 6)
        found = self.store.find_person("+91 98450 00001")
        self.assertEqual([r["data"]["Name"] for r in found], ["DR. ALPHA", "DR. ALPHA"])
        self.assertEqual(len(self.store.find_person("dr. alpha")), 2)
        self.assertEqual(self.store.find_person("nobody@example.com"), [])
        self.assertEqual(self.store.delete_people([r["id"] for r in found]), 2)
        self.assertEqual(self.store.record_count(rid), 4)
        self.assertGreater(self.store.suppression_count(), 0)
        self.assertIsNone(self.store.get_page(self.site.base + "/doc/1"))  # cached copy gone too
        hashes = self.store.db.execute("SELECT value_hash FROM suppressions").fetchall()
        self.assertTrue(all(len(h) == 64 and "98450" not in h for (h,) in hashes))  # hashes only

        events = []
        result = run_recipe(self.store, rid, self.fetcher(), on_progress=events.append)
        self.assertEqual((result.status, result.records), ("done", 0))
        self.assertEqual(self.store.record_count(rid), 4)
        self.assertEqual(self.site.hits["/doc/1"], 1)  # the item page is not fetched again
        skipped = [e.message for e in events if e.event == "skipped"]
        self.assertEqual(skipped, ["Skipped a record on the do-not-collect list."] * 2)

    def test_resume_uses_cache(self):
        rid, _ = self.run_path()
        hits_after_first = dict(self.site.hits)
        second = run_recipe(self.store, rid, self.fetcher())
        self.assertEqual(second.status, "done")
        self.assertEqual(second.records, 0)
        hits_after_second = dict(self.site.hits)
        hits_after_second.pop("/robots.txt", None)
        hits_after_first.pop("/robots.txt", None)
        self.assertEqual(hits_after_first, hits_after_second)  # nothing refetched
        self.assertEqual(self.store.record_count(rid), 6)

    def test_max_records(self):
        _, result = self.run_path(max_records=2)
        self.assertEqual((result.status, result.records), ("done", 2))

    def test_next_link_pagination(self):
        _, result = self.run_path(pagination={"type": "next_link", "selector": "a.next"})
        self.assertEqual((result.status, result.records), ("done", 6))

    def test_block_stops_immediately_without_retry(self):
        rid, result = self.run_path("/blocked-list")
        self.assertEqual(result.status, "blocked")
        self.assertEqual(result.records, 0)
        self.assertEqual(self.site.hits["/limited/1"], 1)
        self.assertEqual(self.site.hits["/limited/2"], 0)
        self.assertEqual(self.store.last_run(rid)["status"], "blocked")

    def test_three_blank_pages_count_as_block(self):
        _, result = self.run_path("/blank-list")
        self.assertEqual(result.status, "blocked")
        self.assertIn("blank pages", result.note)
        self.assertEqual(self.site.hits["/blank/4"], 0)

    def test_robots_disallow(self):
        _, result = self.run_path("/private/list")
        self.assertEqual(result.status, "disallowed")
        self.assertEqual(self.site.hits["/private/list"], 0)

    def test_no_items(self):
        _, result = self.run_path("/empty-list")
        self.assertEqual(result.status, "no_items")

    def test_several_start_pages(self):
        recipe = make_recipe(self.site.base)
        recipe.start_urls = [self.site.base + "/list", self.site.base + "/list-b", self.site.base + "/empty-list"]
        rid = self.store.save_recipe(recipe)
        events = []
        result = run_recipe(self.store, rid, self.fetcher(), on_progress=events.append)
        self.assertEqual((result.status, result.records), ("done", 8))
        starts = {r["data"]["Name"]: r["data"]["_start_page"] for r in self.store.records(rid)}
        self.assertEqual(starts["DR. KAPPA"], self.site.base + "/list-b")
        self.assertEqual(starts["DR. GAMMA"], self.site.base + "/list")
        self.assertEqual(self.site.hits["/list-b"], 2)  # page 1, then page=2 (empty): numbering restarted
        self.assertTrue(any("No items found on this start page" in e.message for e in events))

    def test_max_records_spans_start_pages(self):
        recipe = make_recipe(self.site.base)
        recipe.start_urls = [self.site.base + "/list-b", self.site.base + "/list"]
        rid = self.store.save_recipe(recipe)
        result = run_recipe(self.store, rid, self.fetcher(), max_records=3)
        self.assertEqual((result.status, result.records), ("done", 3))
        self.assertEqual(self.site.hits["/list"], 1)

    def test_stop_control(self):
        control = RunControl()
        control.stop()
        _, result = self.run_path(control=control)
        self.assertEqual((result.status, result.records), ("stopped", 0))


if __name__ == "__main__":
    unittest.main()
